from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


ROOT = Path(__file__).resolve().parents[2]
TEMPLATE_DIR = ROOT / "src" / "frontend" / "public" / "templates"

HEADER_FILL = PatternFill("solid", fgColor="0F766E")
HEADER_FONT = Font(name="Aptos", size=11, bold=True, color="FFFFFF")
HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)

SAMPLE_FILL = PatternFill("solid", fgColor="ECFDF5")
SAMPLE_FONT = Font(name="Aptos", size=11, color="0F172A")
SAMPLE_ALIGNMENT = Alignment(vertical="center", wrap_text=True)

SECTION_FILL = PatternFill("solid", fgColor="CCFBF1")
SECTION_FONT = Font(name="Aptos", size=11, bold=True, color="134E4A")
BODY_FONT = Font(name="Aptos", size=10, color="334155")
TITLE_FONT = Font(name="Aptos Display", size=16, bold=True, color="FFFFFF")
SUBTITLE_FONT = Font(name="Aptos", size=10, italic=True, color="D1FAE5")

THIN_BORDER = Border(
    left=Side(style="thin", color="CBD5E1"),
    right=Side(style="thin", color="CBD5E1"),
    top=Side(style="thin", color="CBD5E1"),
    bottom=Side(style="thin", color="CBD5E1"),
)


@dataclass(frozen=True)
class ColumnSpec:
    header: str
    label: str
    required: bool
    width: int
    help_text: str
    example: str
    kind: str = "text"


@dataclass(frozen=True)
class TemplateSpec:
    filename: str
    title: str
    subtitle: str
    duplicate_rule: str
    columns: tuple[ColumnSpec, ...]
    sample_row: tuple[object, ...]
    extra_notes: tuple[str, ...] = ()


INVOICE_TEMPLATE = TemplateSpec(
    filename="invoice_template.xlsx",
    title="Template import Hoa don",
    subtitle="Ban mau gon de staging parser doc truc tiep tren sheet Data.",
    duplicate_rule=(
        "He thong doi chieu trung theo seller_tax_code + customer_tax_code + "
        "invoice_series + invoice_no + issue_date."
    ),
    columns=(
        ColumnSpec("seller_tax_code", "MST nguoi ban", True, 18, "Nhap dung MST da khai bao tren he thong.", "2300328765"),
        ColumnSpec("customer_tax_code", "MST khach hang", True, 18, "MST ben mua. Neu la chi nhanh, he thong tu doi chieu theo quy tac hoa don.", "0101000002"),
        ColumnSpec("customer_name", "Ten khach hang", False, 28, "Ten ben mua de he thong tao moi customer khi can.", "Cong ty TNHH Mau"),
        ColumnSpec("invoice_template_code", "Ky hieu mau", False, 18, "Ky hieu mau hoa don tren chung tu.", "1C25THK"),
        ColumnSpec("invoice_series", "So hieu hoa don", False, 18, "Series hoa don. Co the de trong neu file nguon khong co.", "AA/26E"),
        ColumnSpec("invoice_no", "So hoa don", True, 16, "So hoa don tren chung tu.", "0000123"),
        ColumnSpec("issue_date", "Ngay phat hanh", True, 16, "Ngay lap hoa don. Ho tro yyyy-MM-dd, dd/MM/yyyy, dd-MM-yyyy.", "2026-01-15", "date"),
        ColumnSpec("revenue_excl_vat", "Doanh so truoc VAT", False, 18, "Gia tri truoc thue. Duoc phep am voi hoa don dieu chinh giam.", "10000000", "money"),
        ColumnSpec("vat_amount", "Tien VAT", False, 16, "Tien thue GTGT. Duoc phep am voi hoa don dieu chinh giam.", "800000", "money"),
        ColumnSpec("total_amount", "Tong thanh toan", False, 18, "Neu de 0 he thong tu tinh = revenue_excl_vat + vat_amount.", "10800000", "money"),
        ColumnSpec("note", "Ghi chu", False, 34, "Dung de ghi dien giai. Neu la hoa don dieu chinh giam, mo ta ro noi dung dieu chinh.", "Hoa don mau cho doi soat"),
    ),
    sample_row=(
        "2300328765",
        "0101000002",
        "Cong ty TNHH Mau",
        "1C25THK",
        "AA/26E",
        "0000123",
        "2026-01-15",
        10_000_000,
        800_000,
        10_800_000,
        "Hoa don mau cho doi soat",
    ),
    extra_notes=(
        "Ngoai template nay, he thong van ho tro ReportDetail.xlsx neu co sheet ExportData.",
        "Khuyen nghi giu nguyen dinh dang text cho cac cot ma so thue va so hoa don.",
    ),
)

ADVANCE_TEMPLATE = TemplateSpec(
    filename="advance_template.xlsx",
    title="Template import Khoan tra ho KH",
    subtitle="So chung tu la bat buoc va duoc dung de doi soat trung tren he thong.",
    duplicate_rule=(
        "He thong doi chieu trung theo seller_tax_code + customer_tax_code + advance_no."
    ),
    columns=(
        ColumnSpec("seller_tax_code", "MST nguoi ban", True, 18, "Nhap dung MST da khai bao tren he thong.", "2300328765"),
        ColumnSpec("customer_tax_code", "MST khach hang", True, 18, "MST ben mua hoac ma thue noi bo dang su dung tren he thong.", "0101000002"),
        ColumnSpec("advance_no", "So chung tu tra ho", True, 18, "So chung tu bat buoc, dung de chong trung voi du lieu da co.", "TH-2026-0001"),
        ColumnSpec("advance_date", "Ngay tra ho", True, 16, "Ngay lap chung tu tra ho. Ho tro yyyy-MM-dd, dd/MM/yyyy, dd-MM-yyyy.", "2026-01-10", "date"),
        ColumnSpec("amount", "So tien", True, 16, "Gia tri tra ho. Phai lon hon 0.", "1500000", "money"),
        ColumnSpec("description", "Dien giai", False, 34, "Mo ta nghiep vu, chi phi, nguon phat sinh.", "Tra ho cuoc van chuyen"),
    ),
    sample_row=(
        "2300328765",
        "0101000002",
        "TH-2026-0001",
        "2026-01-10",
        1_500_000,
        "Tra ho cuoc van chuyen",
    ),
    extra_notes=(
        "Neu trung so chung tu voi ban ghi dang ton tai, dong import se bi canh bao va bo qua.",
    ),
)

RECEIPT_TEMPLATE = TemplateSpec(
    filename="receipt_template.xlsx",
    title="Template import Phieu thu",
    subtitle="Chuan hoa cho doi soat cong no va kiem soat trung so chung tu.",
    duplicate_rule=(
        "He thong doi chieu trung theo seller_tax_code + customer_tax_code + receipt_no."
    ),
    columns=(
        ColumnSpec("seller_tax_code", "MST nguoi ban", True, 18, "Nhap dung MST da khai bao tren he thong.", "2300328765"),
        ColumnSpec("customer_tax_code", "MST khach hang", True, 18, "MST ben mua hoac ma thue noi bo dang su dung tren he thong.", "0101000002"),
        ColumnSpec("receipt_no", "So phieu thu", True, 18, "So chung tu bat buoc, dung de chong trung voi du lieu da co.", "PT-2026-0001"),
        ColumnSpec("receipt_date", "Ngay thu", True, 16, "Ngay lap phieu thu. Ho tro yyyy-MM-dd, dd/MM/yyyy, dd-MM-yyyy.", "2026-01-12", "date"),
        ColumnSpec("applied_period_start", "Ky doi soat", True, 16, "Nhap ngay dau thang cua ky doi soat. Vi du 2026-01-01.", "2026-01-01", "date"),
        ColumnSpec("amount", "So tien", True, 16, "Gia tri thu tien. Phai lon hon 0.", "2500000", "money"),
        ColumnSpec("method", "Phuong thuc", False, 14, "Gia tri hop le: BANK, CASH, OTHER.", "BANK"),
        ColumnSpec("description", "Dien giai", False, 34, "Mo ta nguon thu tien, ky thu, ghi chu doi soat.", "Thu cong no thang 01/2026"),
    ),
    sample_row=(
        "2300328765",
        "0101000002",
        "PT-2026-0001",
        "2026-01-12",
        "2026-01-01",
        2_500_000,
        "BANK",
        "Thu cong no thang 01/2026",
    ),
    extra_notes=(
        "Cot applied_period_start nen la ngay dau thang de he thong phan bo dung ky.",
        "Phuong thuc thu co san dropdown BANK/CASH/OTHER trong cot method.",
    ),
)

TEMPLATES = (INVOICE_TEMPLATE, ADVANCE_TEMPLATE, RECEIPT_TEMPLATE)


def build_data_sheet(spec: TemplateSpec) -> Workbook:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.freeze_panes = "A2"
    sheet.sheet_view.zoomScale = 95
    sheet.sheet_properties.tabColor = "0F766E"

    for index, column in enumerate(spec.columns, start=1):
        cell = sheet.cell(row=1, column=index, value=column.header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGNMENT
        cell.border = THIN_BORDER
        cell.comment = Comment(
            f"{column.label}\n"
            f"{'Bat buoc' if column.required else 'Tuy chon'}\n"
            f"{column.help_text}",
            "Codex",
        )
        sheet.column_dimensions[get_column_letter(index)].width = column.width

    for index, value in enumerate(spec.sample_row, start=1):
        cell = sheet.cell(row=2, column=index, value=value)
        cell.font = SAMPLE_FONT
        cell.fill = SAMPLE_FILL
        cell.alignment = SAMPLE_ALIGNMENT
        cell.border = THIN_BORDER

        kind = spec.columns[index - 1].kind
        if kind == "date":
            cell.number_format = "yyyy-mm-dd"
        elif kind == "money":
            cell.number_format = '#,##0'

    sheet.row_dimensions[1].height = 24
    sheet.row_dimensions[2].height = 22
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(spec.columns))}2"

    if spec.filename == RECEIPT_TEMPLATE.filename:
        method_column = get_column_letter(
            next(index for index, column in enumerate(spec.columns, start=1) if column.header == "method")
        )
        validation = DataValidation(type="list", formula1='"BANK,CASH,OTHER"', allow_blank=True)
        validation.promptTitle = "Phuong thuc thu"
        validation.prompt = "Chon BANK, CASH hoac OTHER."
        sheet.add_data_validation(validation)
        validation.add(f"{method_column}2:{method_column}200")

    return workbook


def add_guide_sheet(workbook: Workbook, spec: TemplateSpec) -> None:
    guide = workbook.create_sheet("HuongDan")
    guide.sheet_view.showGridLines = False
    guide.sheet_view.zoomScale = 90
    guide.sheet_properties.tabColor = "14B8A6"

    for column_letter, width in {
        "A": 24,
        "B": 14,
        "C": 52,
        "D": 28,
        "E": 4,
        "F": 4,
        "G": 4,
        "H": 4,
    }.items():
        guide.column_dimensions[column_letter].width = width

    guide.merge_cells("A1:D1")
    guide["A1"] = spec.title
    guide["A1"].font = TITLE_FONT
    guide["A1"].fill = HEADER_FILL
    guide["A1"].alignment = Alignment(horizontal="left", vertical="center")
    guide.row_dimensions[1].height = 28

    guide.merge_cells("A2:D2")
    guide["A2"] = spec.subtitle
    guide["A2"].font = SUBTITLE_FONT
    guide["A2"].fill = HEADER_FILL
    guide["A2"].alignment = Alignment(horizontal="left", vertical="center")

    write_section_header(guide, 4, "Nguyen tac chung")
    general_notes = [
        "Sheet Data la sheet import chinh. Dong 1 la header co dinh, dong 2 la dong mau de tham chieu.",
        "Khong doi ten header. Co the thay doi thu tu cot, nhung khuyen nghi giu nguyen de doi soat noi bo.",
        "Ngay chap nhan: yyyy-MM-dd, dd/MM/yyyy, dd-MM-yyyy.",
        "Dong co loi se khong duoc ghi so. Dong canh bao co the bi bo qua tuy theo quy tac import.",
    ]
    write_bullets(guide, 5, general_notes)

    write_section_header(guide, 10, "Quy tac chong trung")
    guide.merge_cells("A11:D11")
    guide["A11"] = spec.duplicate_rule
    guide["A11"].font = BODY_FONT
    guide["A11"].alignment = Alignment(wrap_text=True, vertical="top")
    guide["A11"].fill = PatternFill("solid", fgColor="F0FDFA")
    guide["A11"].border = THIN_BORDER

    notes_row = 13
    if spec.extra_notes:
        write_section_header(guide, notes_row, "Luu y nghiep vu")
        write_bullets(guide, notes_row + 1, spec.extra_notes)
        notes_row += len(spec.extra_notes) + 2

    table_header_row = notes_row
    write_section_header(guide, table_header_row, "Mo ta tung cot")

    content_header_row = table_header_row + 1
    for column_index, label in enumerate(("Cot trong file", "Bat buoc", "Y nghia", "Vi du"), start=1):
        cell = guide.cell(row=content_header_row, column=column_index, value=label)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGNMENT
        cell.border = THIN_BORDER

    for row_index, column in enumerate(spec.columns, start=content_header_row + 1):
        row_values = (
            column.header,
            "Co" if column.required else "Khong",
            column.help_text,
            column.example,
        )
        for column_index, value in enumerate(row_values, start=1):
            cell = guide.cell(row=row_index, column=column_index, value=value)
            cell.font = BODY_FONT
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = THIN_BORDER
            if row_index % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F8FAFC")


def write_section_header(sheet, row: int, title: str) -> None:
    sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    cell = sheet.cell(row=row, column=1, value=title)
    cell.font = SECTION_FONT
    cell.fill = SECTION_FILL
    cell.alignment = Alignment(horizontal="left", vertical="center")
    cell.border = THIN_BORDER


def write_bullets(sheet, start_row: int, lines: Iterable[str]) -> None:
    for offset, line in enumerate(lines):
        row = start_row + offset
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        cell = sheet.cell(row=row, column=1, value=f"- {line}")
        cell.font = BODY_FONT
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        cell.border = THIN_BORDER


def save_template(spec: TemplateSpec) -> Path:
    workbook = build_data_sheet(spec)
    add_guide_sheet(workbook, spec)
    TEMPLATE_DIR.mkdir(parents=True, exist_ok=True)
    output_path = TEMPLATE_DIR / spec.filename
    workbook.save(output_path)
    return output_path


def main() -> None:
    for spec in TEMPLATES:
        output_path = save_template(spec)
        print(f"generated {output_path}")


if __name__ == "__main__":
    main()
